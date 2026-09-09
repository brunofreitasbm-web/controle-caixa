"""
Leitor/conciliador das planilhas de meta diaria por loja (Cacau Show).

Formato padrao observado nos exports ("meta loja <codigo> <mes>.xlsx"):
  aba "Export", header na linha 1:
  data | Semana P | $ Venda Bruta | $ Venda Liquida | $ Meta D-1 |
  % Realizado D-1 | $ Meta Total | % Realizado | $ Meta Total Tiquetes |
  $ Venda de Tiquetes | $ Meta Total TM

Este script espelha a logica de processarMetasXLSX() em webapp/app.js
(mesmas colunas candidatas para data/meta, mesma tolerancia de cabecalho
nas primeiras linhas), para permitir ler/conciliar esses arquivos fora do
navegador antes -- ou em vez -- do upload manual em #metas-xlsx-file.

Uso:
  python scratch/importar_metas_xlsx.py "meta loja 4304 setembro.xlsx" ...
  python scratch/importar_metas_xlsx.py *.xlsx --json saida.json
  python scratch/importar_metas_xlsx.py *.xlsx --import --db database.db

Sem --import, o script so le e imprime a conciliacao (dry-run). Com
--import, grava em metas_diarias_lojas (mesmo upsert por loja+data que
POST /api/metas-lojas/importar faz no backend Node).
"""
import argparse
import json
import re
import sqlite3
import sys
from datetime import datetime, date
from pathlib import Path

import openpyxl

LOJA_POR_CODIGO = {
    "9175": "Marambaia",
    "4304": "Icoaraci",
    "9201": "Mário Covas",
}

CANDIDATOS_DATA = ["data", "dia", "data referência", "data referencia"]
CANDIDATOS_META = ["$ meta total", "meta total", "valor meta"]


def normalizar(valor):
    return str(valor).strip().lower() if valor is not None else ""


def inferir_codigo_loja(caminho: Path) -> str:
    m = re.search(r"\b(9175|4304|9201)\b", caminho.stem)
    if not m:
        raise ValueError(
            f"Não foi possível identificar o código da loja no nome do arquivo: {caminho.name}"
        )
    return m.group(1)


def para_data_iso(valor):
    if isinstance(valor, (datetime, date)):
        return valor.strftime("%Y-%m-%d")
    if isinstance(valor, str):
        m = re.match(r"^(\d{1,2})[/-](\d{1,2})[/-](\d{2,4})$", valor.strip())
        if m:
            d, mo, a = m.groups()
            ano = f"20{a}" if len(a) == 2 else a
            return f"{ano}-{mo.zfill(2)}-{d.zfill(2)}"
    return None


def para_valor(valor):
    if isinstance(valor, (int, float)):
        return float(valor)
    texto = re.sub(r"[^\d,.-]", "", str(valor or "0")).replace(",", ".")
    try:
        return float(texto)
    except ValueError:
        return None


def ler_planilha(caminho: Path):
    """Retorna lista de {data, valor, origem} no mesmo formato do payload
    enviado para POST /api/metas-lojas/importar."""
    wb = openpyxl.load_workbook(caminho, data_only=True)
    sheet = wb["Export"] if "Export" in wb.sheetnames else wb.active

    header_row_idx = None
    headers = []
    for r in range(1, min(10, sheet.max_row) + 1):
        row = [sheet.cell(row=r, column=c).value for c in range(1, sheet.max_column + 1)]
        if any("meta total" in normalizar(v) for v in row):
            header_row_idx = r
            headers = row
            break

    if header_row_idx is None:
        raise ValueError(f'Coluna "$ Meta Total" não encontrada em {caminho.name}')

    col_map = {normalizar(h): idx for idx, h in enumerate(headers, start=1) if h}

    def indice_col(candidatos):
        for nome in candidatos:
            if nome in col_map:
                return col_map[nome]
        return None

    idx_data = indice_col(CANDIDATOS_DATA)
    idx_meta = indice_col(CANDIDATOS_META)
    if idx_data is None or idx_meta is None:
        raise ValueError(f'Colunas de Data e/ou "$ Meta Total" ausentes em {caminho.name}')

    linhas = []
    for r in range(header_row_idx + 1, sheet.max_row + 1):
        data_iso = para_data_iso(sheet.cell(row=r, column=idx_data).value)
        if not data_iso:
            continue
        valor = para_valor(sheet.cell(row=r, column=idx_meta).value)
        if valor is None:
            continue
        linhas.append({"data": data_iso, "valor": round(valor, 2), "origem": "diaria"})

    return linhas


def conciliar(arquivos):
    """Le cada arquivo e retorna {loja: [linhas]}, tratando duplicatas de
    data (fica a ultima ocorrencia) como o backend faz."""
    resultado = {}
    for caminho in arquivos:
        codigo = inferir_codigo_loja(caminho)
        loja = LOJA_POR_CODIGO[codigo]
        linhas = ler_planilha(caminho)

        por_data = {}
        for linha in linhas:
            por_data[linha["data"]] = linha
        linhas_dedup = sorted(por_data.values(), key=lambda l: l["data"])

        if loja in resultado:
            print(f"[AVISO] mais de um arquivo para a loja {loja} ({codigo}); mesclando.")
            existentes = {l["data"]: l for l in resultado[loja]}
            existentes.update(por_data)
            linhas_dedup = sorted(existentes.values(), key=lambda l: l["data"])

        resultado[loja] = linhas_dedup
    return resultado


def imprimir_relatorio(resultado):
    for loja, linhas in resultado.items():
        total = sum(l["valor"] for l in linhas)
        print(f"\n=== {loja} — {len(linhas)} dia(s) ===")
        if linhas:
            print(f"  período: {linhas[0]['data']} a {linhas[-1]['data']}")
        print(f"  soma das metas diárias: {total:,.2f}")
        for l in linhas:
            print(f"  {l['data']}  R$ {l['valor']:,.2f}")


def importar_para_db(resultado, caminho_db: str):
    conn = sqlite3.connect(caminho_db)
    criado_em = datetime.now().isoformat()
    total_gravado = 0
    for loja, linhas in resultado.items():
        for linha in linhas:
            id_registro = f"{loja}_{linha['data']}"
            conn.execute(
                """
                INSERT INTO metas_diarias_lojas (id, loja, data, valor, origem, criadoEm)
                VALUES (?, ?, ?, ?, ?, ?)
                ON CONFLICT(loja, data) DO UPDATE SET
                    valor = excluded.valor,
                    origem = excluded.origem
                """,
                (id_registro, loja, linha["data"], linha["valor"], linha["origem"], criado_em),
            )
            total_gravado += 1
    conn.commit()
    conn.close()
    print(f"\n{total_gravado} registro(s) gravado(s) em {caminho_db} (tabela metas_diarias_lojas).")


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("arquivos", nargs="+", help="Planilhas .xlsx de meta por loja")
    parser.add_argument("--json", help="Grava a conciliação em um arquivo JSON (formato do payload de /api/metas-lojas/importar)")
    parser.add_argument("--import", dest="importar", action="store_true", help="Grava direto no SQLite (--db)")
    parser.add_argument("--db", default="database.db", help="Caminho do SQLite (default: database.db)")
    args = parser.parse_args()

    arquivos = [Path(a) for a in args.arquivos]
    for a in arquivos:
        if not a.exists():
            sys.exit(f"Arquivo não encontrado: {a}")

    resultado = conciliar(arquivos)
    imprimir_relatorio(resultado)

    if args.json:
        payload = [
            {"loja": loja, "linhas": linhas}
            for loja, linhas in resultado.items()
        ]
        Path(args.json).write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
        print(f"\nJSON salvo em {args.json}")

    if args.importar:
        importar_para_db(resultado, args.db)


if __name__ == "__main__":
    main()
